"""Exact reproduction of DESIGN.md section 15 worked example."""

from __future__ import annotations

import csv
from pathlib import Path

import openpyxl
import pytest

from comparator.runner import process_pair
from comparator.files import FilePair
from comparator.settings import Settings
from tests.make_test_data import write_worked_example, ROOT, DS, INFA


def _summary_map(ws) -> dict:
    """Map first-column labels to row values."""
    out = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        key = str(row[0])
        out[key] = row
    return out


@pytest.fixture
def customer_settings(tmp_path) -> Settings:
    write_worked_example()
    out = tmp_path / "OUTPUT"
    temp = out / "_tmp"
    out.mkdir()
    temp.mkdir()
    return Settings(
        ds_folder=DS,
        infa_folder=INFA,
        output_folder=out,
        temp_folder=temp,
        excel_sample_rows=1000,
    )


def test_worked_example_with_key(customer_settings, tmp_path):
    pair = FilePair(
        prefix="customer",
        ds_path=DS / "customer_ds.csv",
        infa_path=INFA / "customer_infa.csv",
    )
    from comparator.runlog import RunLog

    result = process_pair(pair, customer_settings, RunLog(customer_settings.output_folder))

    assert result.overall == "MISMATCH"
    assert result.columns_match is True
    assert result.ds.row_count == 7
    assert result.infa.row_count == 6
    assert result.ds.distinct_rows == 6
    assert result.infa.distinct_rows == 6
    assert result.row_compare.matched_rows == 2
    assert result.row_compare.only_ds == 4
    assert result.row_compare.only_infa == 4
    assert result.row_compare.extra_ds == 1
    assert result.row_compare.extra_infa == 0
    assert abs(result.row_compare.match_rate - 28.57) < 0.01

    assert result.key is not None
    assert result.key.source == "Auto-detected"
    assert result.key.columns == ["CUST_ID"]
    assert result.key.paired_records == 3
    assert result.key.differing_cells == 5
    assert result.key.unpaired_by_status.get("Only in DataStage") == 1
    assert result.key.unpaired_by_status.get("Only in Informatica") == 1
    assert result.key.unpaired_by_status.get("Extra copies in DataStage") == 1

    # Column differences
    by_name = {c.name: c for c in result.column_diffs}
    assert by_name["CUST_ID"].paired_differing == "Key"
    assert by_name["NAME"].only_ds == 3
    assert by_name["NAME"].only_infa == 3
    assert by_name["NAME"].paired_differing == 2
    assert by_name["AMOUNT"].paired_differing == 2
    assert by_name["CITY"].paired_differing == 1
    assert by_name["CITY"].ds_examples == ["Nashik"]
    assert by_name["CITY"].infa_examples == ["Pune "]

    xlsx = customer_settings.output_folder / "customer_comparison.xlsx"
    assert xlsx.exists()
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    # Differences sheet order
    diff_ws = wb["Differences (sample)"]
    diff_rows = list(diff_ws.iter_rows(values_only=True))
    assert diff_rows[0][:6] == (
        "CUST_ID",
        "Column",
        "INFA",
        "DS",
        "INFA copies",
        "DS copies",
    )
    data = diff_rows[1:]
    expected_diff = [
        ("101", "AMOUNT", "'0000000010.00'", "'10.00'", 1, 1),
        ("104", "NAME", "'Meera'", "'meera'", 1, 1),
        ("106", "NAME", "'Rohit'", "'rohit'", 1, 1),
        ("106", "AMOUNT", "'00010.000'", "'10.000'", 1, 1),
        ("106", "CITY", "'Pune '", "'Pune'", 1, 1),
    ]
    assert len(data) == 5
    for got, exp in zip(data, expected_diff):
        got6 = list(got[:6])
        for i in (4, 5):
            if isinstance(got6[i], str) and got6[i].isdigit():
                got6[i] = int(got6[i])
        assert tuple(got6) == exp

    # Unpaired sheet
    un_ws = wb["Unpaired rows (sample)"]
    un_rows = list(un_ws.iter_rows(values_only=True))[1:]
    statuses = {r[0]: r for r in un_rows}
    assert "Extra copies in DataStage" in statuses
    assert statuses["Extra copies in DataStage"][4:8] == ("103", "ravi", "5.00", "Delhi")
    assert statuses["Only in Informatica"][4:8] == ("105", "kiran", "3.00", "Pune")
    assert statuses["Only in DataStage"][4:8] == ("107", "sana", "1.50", "Nashik")

    # Differences CSV as a set (order not guaranteed)
    diff_csv = customer_settings.output_folder / "customer_comparison_differences.csv"
    assert diff_csv.exists()
    with diff_csv.open(encoding="utf-8", newline="") as f:
        reader = csv.reader(f)
        header = next(reader)
        assert header == [
            "CUST_ID",
            "Column",
            "INFA",
            "DS",
            "INFA copies",
            "DS copies",
        ]
        rows = {tuple(r) for r in reader}
    expected_csv = {
        ("101", "AMOUNT", "0000000010.00", "10.00", "1", "1"),
        ("104", "NAME", "Meera", "meera", "1", "1"),
        ("106", "NAME", "Rohit", "rohit", "1", "1"),
        ("106", "AMOUNT", "00010.000", "10.000", "1", "1"),
        ("106", "CITY", "Pune ", "Pune", "1", "1"),
    }
    assert rows == expected_csv

    # Trailing space on last CITY INFA value
    city_row = [r for r in rows if r[1] == "CITY"][0]
    assert city_row[2] == "Pune "
    assert city_row[2].endswith(" ")


def test_worked_example_fallback_no_key(tmp_path):
    write_worked_example()
    out = tmp_path / "OUTPUT"
    temp = out / "_tmp"
    out.mkdir()
    temp.mkdir()
    settings = Settings(
        ds_folder=DS,
        infa_folder=INFA,
        output_folder=out,
        temp_folder=temp,
        key_overrides={"customer": None},
        excel_sample_rows=1000,
    )
    pair = FilePair(
        prefix="customer",
        ds_path=DS / "customer_ds.csv",
        infa_path=INFA / "customer_infa.csv",
    )
    from comparator.runlog import RunLog

    result = process_pair(pair, settings, RunLog(out))
    assert result.key is not None
    assert result.key.source == "None"
    assert result.mismatches_table is not None

    xlsx = out / "customer_comparison.xlsx"
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["Mismatches (sample)"]
    rows = list(ws.iter_rows(values_only=True))[1:]
    # Expected order: sorted by all columns then status
    expected = [
        ("Only in Informatica", 0, 1, 1, "101", "nigam", "0000000010.00", "Pune"),
        ("Only in DataStage", 1, 0, 1, "101", "nigam", "10.00", "Pune"),
        ("Extra copies in DataStage", 2, 1, 1, "103", "ravi", "5.00", "Delhi"),
        ("Only in Informatica", 0, 1, 1, "104", "Meera", "7.25", "Goa"),
        ("Only in DataStage", 1, 0, 1, "104", "meera", "7.25", "Goa"),
        ("Only in Informatica", 0, 1, 1, "105", "kiran", "3.00", "Pune"),
        ("Only in Informatica", 0, 1, 1, "106", "Rohit", "00010.000", "Pune "),
        ("Only in DataStage", 1, 0, 1, "106", "rohit", "10.000", "Pune"),
        ("Only in DataStage", 1, 0, 1, "107", "sana", "1.50", "Nashik"),
    ]
    assert len(rows) == 9
    for got, exp in zip(rows, expected):
        got_n = list(got[:8])
        for i in (1, 2, 3):
            if isinstance(got_n[i], str) and got_n[i].lstrip("-").isdigit():
                got_n[i] = int(got_n[i])
        assert tuple(got_n) == exp
